#!/usr/bin/env python3
"""Convert the local-only formal test account CSV into a formatted XLSX workbook."""

from __future__ import annotations

import argparse
import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = [
    "account_type",
    "needo_id",
    "nickname",
    "email",
    "password",
]
HEADER_LABELS = [
    "账号类型",
    "NeeDoID",
    "昵称",
    "邮件",
    "密码",
]
TYPE_LABELS = {
    "merchant_owner": "店铺服务号",
    "technician": "技师",
    "customer": "一般用户",
    "admin": "运营后台超级管理员",
    "operator": "运营人员",
    "platform_admin": "平台管理员",
    "broker": "渠道合作方",
    "affiliate": "渠道合作方",
}
NAVY = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F3F6F9"
WHITE = "FFFFFF"
TEXT = "1F2937"
WORKBOOK_FONT = "Hiragino Sans GB"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("csv_path", type=Path)
    parser.add_argument("xlsx_path", type=Path)
    return parser.parse_args()


def load_rows(csv_path: Path) -> list[dict[str, str]]:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if reader.fieldnames != HEADERS:
            raise ValueError(f"unexpected account CSV headers: {reader.fieldnames}")
        rows = list(reader)
    if len(rows) != 216:
        raise ValueError(f"expected 216 account rows, found {len(rows)}")
    return rows


def create_workbook(rows: list[dict[str, str]], output_path: Path) -> None:
    workbook = Workbook()
    accounts = workbook.active
    accounts.title = "Accounts"
    accounts.freeze_panes = "A2"
    accounts.sheet_view.showGridLines = False
    accounts.sheet_properties.pageSetUpPr.fitToPage = True
    accounts.page_setup.orientation = "landscape"
    accounts.page_setup.fitToWidth = 1
    accounts.page_setup.fitToHeight = 0
    accounts.print_title_rows = "1:1"
    accounts.append(HEADER_LABELS)
    for row in rows:
        accounts.append(
            [
                TYPE_LABELS.get(row["account_type"], row["account_type"]),
                row["needo_id"],
                row["nickname"],
                row["email"],
                row["password"],
            ]
        )

    thin = Side(style="thin", color="D5DCE5")
    for cell in accounts[1]:
        cell.font = Font(name=WORKBOOK_FONT, size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    accounts.row_dimensions[1].height = 24

    for row_index in range(2, accounts.max_row + 1):
        for column_index in range(1, accounts.max_column + 1):
            cell = accounts.cell(row=row_index, column=column_index)
            cell.font = Font(name=WORKBOOK_FONT, size=10, color=TEXT)
            cell.alignment = Alignment(vertical="center")
            cell.number_format = "@"
        accounts.row_dimensions[row_index].height = 20

    widths = {"A": 20, "B": 16, "C": 32, "D": 38, "E": 18}
    for column, width in widths.items():
        accounts.column_dimensions[column].width = width

    table = Table(displayName="FormalTestAccounts", ref=f"A1:E{accounts.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    accounts.add_table(table)

    summary = workbook.create_sheet("Read Me")
    summary.sheet_view.showGridLines = False
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 58
    summary.append(["项目", "说明"])
    role_counts = Counter(row["account_type"] for row in rows)
    summary_rows = [
        ("用途", "NeeDo 正式本地测试账号，用于真实 API 和三端页面验收"),
        ("数据期间", "2026-06-01 至 2026-08-31"),
        ("店铺服务号", str(role_counts["merchant_owner"] + role_counts["店铺服务号"])),
        ("技师账号", str(role_counts["technician"] + role_counts["技师"])),
        ("一般用户", str(role_counts["customer"] + role_counts["一般用户"])),
        (
            "平台及运营账号",
            str(
                role_counts["admin"]
                + role_counts["platform_admin"]
                + role_counts["operator"]
                + role_counts["运营后台超级管理员"]
                + role_counts["运营后台运营管理员"]
            ),
        ),
        (
            "渠道合作账号",
            str(
                role_counts["affiliate"]
                + role_counts["broker"]
                + role_counts["渠道合作账号"]
            ),
        ),
        ("账号总数", str(len(rows))),
        ("动态内容", "每账号 15 条正式动态：纯文字、单图、多图、视频、引用各 3 条。"),
        ("好友关系", "每账号 36 个双向好友，同时包含店铺服务号、技师和一般用户。"),
        ("安全提示", "仅限本地测试。包含统一测试密码，不得提交、公开或用于生产环境。"),
        ("统一密码", f"全部账号使用用户指定的测试密码：{rows[0]['password']}"),
    ]
    for item in summary_rows:
        summary.append(item)
    for cell in summary[1]:
        cell.font = Font(name=WORKBOOK_FONT, size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index in range(2, summary.max_row + 1):
        summary.cell(row=row_index, column=1).font = Font(
            name=WORKBOOK_FONT, size=10, bold=True, color=TEXT
        )
        summary.cell(row=row_index, column=1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        summary.cell(row=row_index, column=2).font = Font(
            name=WORKBOOK_FONT, size=10, color=TEXT
        )
        summary.cell(row=row_index, column=2).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        for column_index in (1, 2):
            summary.cell(row=row_index, column=column_index).alignment = Alignment(
                vertical="top", wrap_text=True
            )
            summary.cell(row=row_index, column=column_index).border = Border(bottom=thin)
        summary.row_dimensions[row_index].height = 30
    summary.freeze_panes = "A2"

    workbook.properties.title = "NeeDo Formal Test Accounts"
    workbook.properties.subject = "Local formal test login accounts"
    workbook.properties.creator = "NeeDo"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def verify_workbook(output_path: Path) -> None:
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    if workbook.sheetnames != ["Accounts", "Read Me"]:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")
    accounts = workbook["Accounts"]
    if accounts.max_row != 217 or accounts.max_column != 5:
        raise ValueError(
            f"unexpected account sheet shape: {accounts.max_row} rows x {accounts.max_column} columns"
        )
    if [accounts.cell(1, column).value for column in range(1, 6)] != HEADER_LABELS:
        raise ValueError("account workbook headers do not match the approved export contract")


def main() -> None:
    args = parse_args()
    rows = load_rows(args.csv_path)
    create_workbook(rows, args.xlsx_path)
    verify_workbook(args.xlsx_path)
    print(
        f'{{"xlsx":"{args.xlsx_path.resolve()}","accountRows":{len(rows)},"status":"ok"}}'
    )


if __name__ == "__main__":
    main()
