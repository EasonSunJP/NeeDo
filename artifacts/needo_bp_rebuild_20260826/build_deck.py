from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE = Path(__file__).resolve().parent
ORIGINAL = Path("/Users/eason/Documents/NeeDo/NeeDo_BP.pptx")
REVISED = Path("/Users/eason/Documents/NeeDo/NeeDo_BP_投资人精简优化版.pptx")
OUTPUT = BASE / "NeeDo_BP_投资人精简优化_美术重制版.pptx"
REPORT = BASE / "qa_report.txt"

# Revised slide number -> original slide number.
SLIDE_MAP = [1, 2, 3, 4, 7, 10, 11, 9, 21, 22, 20, 19, 18, 23, 28, 25, 26, 32, 33, 34]


def shape_key(shape: Any) -> tuple[str, int, int, int, int]:
    return (shape.name, shape.left, shape.top, shape.width, shape.height)


def has_transferable_text(shape: Any) -> bool:
    return bool(getattr(shape, "has_text_frame", False) or getattr(shape, "has_table", False))


def copy_text_content(source: Any, target: Any) -> None:
    if getattr(source, "has_text_frame", False) and getattr(target, "has_text_frame", False):
        source_body = source.text_frame._txBody
        target_body = target.text_frame._txBody
        target_body.getparent().replace(target_body, deepcopy(source_body))
        return

    if getattr(source, "has_table", False) and getattr(target, "has_table", False):
        if len(source.table.rows) != len(target.table.rows) or len(source.table.columns) != len(target.table.columns):
            raise ValueError(f"Table dimensions differ for {source.name}")
        for source_row, target_row in zip(source.table.rows, target.table.rows):
            for source_cell, target_cell in zip(source_row.cells, target_row.cells):
                source_body = source_cell._tc.txBody
                target_body = target_cell._tc.txBody
                target_cell._tc.replace(target_body, deepcopy(source_body))
        return

    raise ValueError(f"Text-capability mismatch for {source.name}")


def transfer_revised_text(source_slide: Any, target_slide: Any) -> tuple[int, list[str]]:
    targets = {shape_key(shape): shape for shape in target_slide.shapes}
    matched = 0
    added: list[str] = []

    for source_shape in source_slide.shapes:
        if not has_transferable_text(source_shape):
            continue
        target_shape = targets.get(shape_key(source_shape))
        if target_shape is not None:
            copy_text_content(source_shape, target_shape)
            matched += 1
            continue

        text = ""
        if getattr(source_shape, "has_text_frame", False):
            text = source_shape.text.strip()
        elif getattr(source_shape, "has_table", False):
            text = " ".join(cell.text for row in source_shape.table.rows for cell in row.cells).strip()
        if text:
            target_slide.shapes._spTree.insert_element_before(deepcopy(source_shape._element), "p:extLst")
            added.append(source_shape.name)

    return matched, added


def replace_single_run_text(shape: Any, value: str) -> None:
    paragraph = shape.text_frame.paragraphs[0]
    if not paragraph.runs:
        paragraph.add_run().text = value
    else:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    for extra in shape.text_frame.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def style_run(run: Any, *, size: float, bold: bool, color: str) -> None:
    run.font.name = "MiSans Normal"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_bounds(shape: Any, *, left: float, top: float, width: float, height: float) -> None:
    shape.left = Inches(left)
    shape.top = Inches(top)
    shape.width = Inches(width)
    shape.height = Inches(height)


def format_text_box(shape: Any, *, size: float, bold: bool, color: str, align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in frame.paragraphs:
        paragraph.alignment = align
        paragraph.space_before = 0
        paragraph.space_after = 0
        paragraph.line_spacing = 1.0
        for run in paragraph.runs:
            style_run(run, size=size, bold=bold, color=color)


def apply_layout_fixes(slide_number: int, slide: Any) -> None:
    bottom_backgrounds = [
        shape for shape in slide.shapes if shape.name == "Rounded card" and shape.top > Inches(5.5)
    ]
    bottom_bodies = [
        shape for shape in slide.shapes if shape.name == "Card body" and shape.top > Inches(5.5)
    ]

    if slide_number == 5:
        background = bottom_backgrounds[0]
        body = bottom_bodies[0]
        insight = next(shape for shape in slide.shapes if shape.name == "TextBox 24")
        set_bounds(background, left=0.55, top=5.74, width=12.2, height=0.75)
        set_bounds(insight, left=0.77, top=5.84, width=11.76, height=0.27)
        set_bounds(body, left=0.77, top=6.20, width=11.76, height=0.22)
        format_text_box(insight, size=11.5, bold=True, color="0B5943")
        return

    if slide_number == 7:
        background = bottom_backgrounds[0]
        body = bottom_bodies[0]
        target = next(shape for shape in slide.shapes if shape.name == "TextBox 19")
        set_bounds(background, left=0.55, top=5.74, width=12.2, height=0.75)
        set_bounds(body, left=0.77, top=5.86, width=11.76, height=0.22)
        set_bounds(target, left=0.77, top=6.20, width=11.76, height=0.22)
        format_text_box(target, size=11.5, bold=True, color="0B5943", align=PP_ALIGN.CENTER)
        return

    if slide_number == 10:
        background = bottom_backgrounds[0]
        body = bottom_bodies[0]
        set_bounds(background, left=0.55, top=5.72, width=12.2, height=0.84)
        set_bounds(body, left=0.77, top=5.84, width=11.76, height=0.60)
        format_text_box(body, size=10.5, bold=False, color="101A16")


def set_team_body(shape: Any, name: str, description: str) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP

    name_paragraph = frame.paragraphs[0]
    name_paragraph.alignment = PP_ALIGN.LEFT
    name_paragraph.space_after = Pt(12)
    name_run = name_paragraph.add_run()
    name_run.text = name
    style_run(name_run, size=13.5, bold=True, color="0B5943")

    description_paragraph = frame.add_paragraph()
    description_paragraph.alignment = PP_ALIGN.LEFT
    description_paragraph.line_spacing = 1.08
    description_run = description_paragraph.add_run()
    description_run.text = description
    style_run(description_run, size=10.5, bold=False, color="101A16")


def rebuild_team_slide(revised_slide: Any, target_slide: Any) -> None:
    target_by_name = {shape.name: shape for shape in target_slide.shapes}
    replace_single_run_text(target_by_name["Page number"], "19")
    replace_single_run_text(target_by_name["Slide title"], "核心團隊")

    team_boxes = [
        shape
        for shape in revised_slide.shapes
        if shape.name.startswith("TextBox") and getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]
    team_boxes.sort(key=lambda shape: (shape.left, shape.top))
    if len(team_boxes) != 12:
        raise ValueError(f"Expected 12 team text boxes, found {len(team_boxes)}")
    people = [tuple(team_boxes[index : index + 3]) for index in range(0, 12, 3)]

    card_titles = sorted(
        [shape for shape in target_slide.shapes if shape.name == "Card title"],
        key=lambda shape: shape.left,
    )
    card_bodies = sorted(
        [shape for shape in target_slide.shapes if shape.name == "Card body" and shape.top < 5 * 914400],
        key=lambda shape: shape.left,
    )
    if len(card_titles) != 4 or len(card_bodies) != 4:
        raise ValueError("The team template does not contain four title/body card pairs")

    for title_shape, body_shape, (role_shape, name_shape, description_shape) in zip(
        card_titles, card_bodies, people
    ):
        replace_single_run_text(title_shape, role_shape.text.strip())
        set_team_body(body_shape, name_shape.text.strip(), description_shape.text.strip())

    for shape in list(target_slide.shapes):
        if shape.top > 5.8 * 914400:
            target_slide.shapes._spTree.remove(shape._element)


def retain_and_reorder_slides(prs: Presentation, slide_numbers: list[int]) -> None:
    slide_ids = list(prs.slides._sldIdLst)
    selected = [slide_ids[number - 1] for number in slide_numbers]
    selected_rel_ids = {slide_id.rId for slide_id in selected}

    for slide_id in slide_ids:
        if slide_id.rId not in selected_rel_ids:
            prs.part.drop_rel(slide_id.rId)

    prs.slides._sldIdLst.clear()
    for slide_id in selected:
        prs.slides._sldIdLst.append(slide_id)


def main() -> None:
    original = Presentation(ORIGINAL)
    revised = Presentation(REVISED)
    if (original.slide_width, original.slide_height) != (revised.slide_width, revised.slide_height):
        raise ValueError("Source deck slide sizes differ")
    if len(revised.slides) != 20:
        raise ValueError(f"Expected 20 revised slides, found {len(revised.slides)}")

    retain_and_reorder_slides(original, SLIDE_MAP)
    if len(original.slides) != 20:
        raise ValueError(f"Expected 20 retained slides, found {len(original.slides)}")

    report_lines = ["NeeDo BP rebuild report", "", f"Original: {ORIGINAL}", f"Revised: {REVISED}"]
    for index, (revised_slide, target_slide) in enumerate(zip(revised.slides, original.slides), start=1):
        if index == 19:
            rebuild_team_slide(revised_slide, target_slide)
            report_lines.append("Slide 19: rebuilt with original four-card design")
            continue
        matched, added = transfer_revised_text(revised_slide, target_slide)
        apply_layout_fixes(index, target_slide)
        suffix = f"; added {', '.join(added)}" if added else ""
        report_lines.append(f"Slide {index:02d}: transferred {matched} text/table shapes{suffix}")

    original.core_properties.title = "NeeDo 海外投資人 Pre-A 融資簡報｜投資人精簡優化美術重製版"
    original.core_properties.subject = "20-page investor edition rebuilt with original NeeDo art direction"
    original.core_properties.comments = "Generated from two user-provided source decks without overwriting either source."
    original.save(OUTPUT)
    report_lines.extend(["", f"Output: {OUTPUT}", f"Slides: {len(original.slides)}"])
    REPORT.write_text("\n".join(report_lines) + "\n", encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    main()
