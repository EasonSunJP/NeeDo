from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


BASE = Path(__file__).resolve().parent
REVISED = Path("/Users/eason/Documents/NeeDo/NeeDo_BP_投资人精简优化版.pptx")
OUTPUT = BASE / "NeeDo_BP_投资人精简优化_美术重制版.pptx"
REPORT = BASE / "qa_report.txt"


def normalized(value: str) -> str:
    return " ".join(value.replace("\u3000", " ").split())


def slide_texts(slide: Any) -> list[str]:
    texts: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            value = normalized(shape.text)
            if value:
                texts.append(value)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    value = normalized(cell.text)
                    if value:
                        texts.append(value)
    return texts


def slide_title(slide: Any) -> str:
    for shape in slide.shapes:
        if shape.name in {"Slide title", "Text 2", "Text 3"} and getattr(shape, "has_text_frame", False):
            value = normalized(shape.text)
            if value:
                return value
    return ""


def count_type(slide: Any, shape_type: MSO_SHAPE_TYPE) -> int:
    return sum(shape.shape_type == shape_type for shape in slide.shapes)


def main() -> None:
    revised = Presentation(REVISED)
    output = Presentation(OUTPUT)
    failures: list[str] = []
    lines = ["", "Automated content QA"]

    if len(output.slides) != 20:
        failures.append(f"Output has {len(output.slides)} slides instead of 20")

    for index, (revised_slide, output_slide) in enumerate(zip(revised.slides, output.slides), start=1):
        title_expected = slide_title(revised_slide)
        title_actual = slide_title(output_slide)
        if title_actual != title_expected:
            failures.append(f"Slide {index:02d} title mismatch: {title_actual!r} != {title_expected!r}")

        chart_expected = count_type(revised_slide, MSO_SHAPE_TYPE.CHART)
        chart_actual = count_type(output_slide, MSO_SHAPE_TYPE.CHART)
        if chart_actual != chart_expected:
            failures.append(f"Slide {index:02d} chart count mismatch: {chart_actual} != {chart_expected}")

        if index != 19:
            expected = Counter(slide_texts(revised_slide))
            actual = Counter(slide_texts(output_slide))
            if actual != expected:
                missing = list((expected - actual).elements())
                extra = list((actual - expected).elements())
                failures.append(f"Slide {index:02d} text mismatch; missing={missing}; extra={extra}")
        else:
            actual_values = set(slide_texts(output_slide))
            actual_joined = "\n".join(actual_values)
            required = {
                "核心團隊",
                "Founder",
                "CEO",
                "CTO",
                "CPO",
                "姓名待補",
                "Vision / Capital / Strategic Partnerships",
                "Business Execution / Fundraising / Organization",
                "Architecture / Platform / AI & Data / Security",
                "Product / UX / Marketplace / SaaS / Localization",
            }
            missing = sorted(value for value in required if value not in actual_joined)
            if missing:
                failures.append(f"Slide 19 missing team content: {missing}")
            timeline_residue = {f"{number:02d}" for number in range(1, 7)} & actual_values
            if timeline_residue:
                failures.append(f"Slide 19 retains timeline labels: {sorted(timeline_residue)}")

        for shape in output_slide.shapes:
            if shape.left < 0 or shape.top < 0:
                failures.append(f"Slide {index:02d} shape {shape.name!r} starts outside the slide")
            if shape.left + shape.width > output.slide_width or shape.top + shape.height > output.slide_height:
                failures.append(f"Slide {index:02d} shape {shape.name!r} exceeds slide bounds")

        lines.append(
            f"Slide {index:02d}: title OK; text {'custom-team check' if index == 19 else 'exact'}; "
            f"charts={chart_actual}; pictures={count_type(output_slide, MSO_SHAPE_TYPE.PICTURE)}"
        )

    page_numbers = []
    for slide in output.slides:
        number_shapes = [
            normalized(shape.text)
            for shape in slide.shapes
            if shape.name == "Page number" and getattr(shape, "has_text_frame", False)
        ]
        page_numbers.append(number_shapes[0] if number_shapes else "")
    expected_numbers = [f"{number:02d}" for number in range(1, 21)]
    if page_numbers != expected_numbers:
        failures.append(f"Page numbers mismatch: {page_numbers}")

    lines.extend(["", f"Page numbers: {', '.join(page_numbers)}"])
    if failures:
        lines.extend(["", "FAILURES", *failures])
        REPORT.write_text(REPORT.read_text(encoding="utf-8") + "\n".join(lines) + "\n", encoding="utf-8")
        raise SystemExit("\n".join(failures))

    lines.extend(["", "Automated content QA: PASS"])
    REPORT.write_text(REPORT.read_text(encoding="utf-8") + "\n".join(lines) + "\n", encoding="utf-8")
    print("Automated content QA: PASS")


if __name__ == "__main__":
    main()
