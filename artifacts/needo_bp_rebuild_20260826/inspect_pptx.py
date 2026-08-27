from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


BASE = Path(__file__).resolve().parent
ORIGINAL = Path("/Users/eason/Documents/NeeDo/NeeDo_BP.pptx")
REVISED = Path("/Users/eason/Documents/NeeDo/NeeDo_BP_投资人精简优化版.pptx")


def shape_record(shape: Any) -> dict[str, Any]:
    record: dict[str, Any] = {
        "name": shape.name,
        "type": str(shape.shape_type),
        "left": round(shape.left / 914400, 3),
        "top": round(shape.top / 914400, 3),
        "width": round(shape.width / 914400, 3),
        "height": round(shape.height / 914400, 3),
    }
    if getattr(shape, "has_text_frame", False):
        record["text"] = " | ".join(p.text for p in shape.text_frame.paragraphs).strip()
    elif getattr(shape, "has_table", False):
        record["text"] = " || ".join(
            " | ".join(cell.text for cell in row.cells) for row in shape.table.rows
        )
    elif shape.shape_type == MSO_SHAPE_TYPE.CHART:
        record["chart"] = True
    return record


def deck_record(path: Path) -> dict[str, Any]:
    prs = Presentation(path)
    return {
        "path": str(path),
        "width": prs.slide_width,
        "height": prs.slide_height,
        "slides": [
            {
                "number": index + 1,
                "shapes": [shape_record(shape) for shape in slide.shapes],
            }
            for index, slide in enumerate(prs.slides)
        ],
    }


if __name__ == "__main__":
    output = BASE / "shape_inventory.json"
    output.write_text(
        json.dumps({"original": deck_record(ORIGINAL), "revised": deck_record(REVISED)}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(output)
