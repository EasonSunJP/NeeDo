#!/usr/bin/env python3
"""Read-only extractor for the NeeDo 36-month financial model."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


ROWS = {
    "stores": 20,
    "technicians": 23,
    "orders": 34,
    "revenue": 50,
    "opex": 72,
    "operating_profit": 74,
    "operating_margin": 75,
    "ending_cash": 80,
}


def row_values(ws, row: int) -> list[float]:
    values: list[float] = []
    for col in range(3, 39):
        value = ws.cell(row=row, column=col).value
        values.append(float(value or 0))
    return values


def scenario(path: Path, sheet_name: str) -> dict[str, object]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    monthly = {key: row_values(ws, row) for key, row in ROWS.items()}
    annual = []
    for year in range(3):
        start = year * 12
        end = start + 12
        annual.append(
            {
                "year": year + 1,
                "stores": monthly["stores"][end - 1],
                "technicians": monthly["technicians"][end - 1],
                "orders": sum(monthly["orders"][start:end]),
                "revenue": sum(monthly["revenue"][start:end]),
                "opex": sum(monthly["opex"][start:end]),
                "operating_profit": sum(monthly["operating_profit"][start:end]),
                "operating_margin": (
                    sum(monthly["operating_profit"][start:end])
                    / sum(monthly["revenue"][start:end])
                    if sum(monthly["revenue"][start:end])
                    else 0
                ),
                "ending_cash": monthly["ending_cash"][end - 1],
            }
        )
    positive_month = next(
        (
            index + 1
            for index, value in enumerate(monthly["operating_profit"])
            if value > 0 and all(v > 0 for v in monthly["operating_profit"][index:])
        ),
        None,
    )
    minimum_cash = min(monthly["ending_cash"])
    minimum_cash_month = monthly["ending_cash"].index(minimum_cash) + 1
    return {
        "source": str(path),
        "sheet": sheet_name,
        "monthly": monthly,
        "annual": annual,
        "positive_month": positive_month,
        "minimum_cash": minimum_cash,
        "minimum_cash_month": minimum_cash_month,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base", type=Path, required=True)
    parser.add_argument("--aggressive", type=Path, required=True)
    args = parser.parse_args()
    payload = {
        "conservative": scenario(args.base, "A_36个月模型"),
        "general": scenario(args.base, "B_36个月模型"),
        "aggressive": scenario(args.aggressive, "B_36个月模型"),
    }
    print(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    main()
