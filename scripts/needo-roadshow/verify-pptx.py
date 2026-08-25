#!/usr/bin/env python3
"""Artifact-level verification for the NeeDo investor deck."""

from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

from pptx import Presentation


EMU_PER_INCH = 914400


def shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False):
        return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
    return ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()

    errors: list[str] = []
    warnings: list[str] = []
    with zipfile.ZipFile(args.pptx) as archive:
        bad = archive.testzip()
        names = set(archive.namelist())
        if bad:
            errors.append(f"ZIP CRC failure: {bad}")
        for required in ("[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"):
            if required not in names:
                errors.append(f"Missing package part: {required}")

    deck = Presentation(args.pptx)
    width = deck.slide_width / EMU_PER_INCH
    height = deck.slide_height / EMU_PER_INCH
    if len(deck.slides) != 34:
        errors.append(f"Expected 34 slides, found {len(deck.slides)}")
    if abs(width - 13.333) > 0.02 or abs(height - 7.5) > 0.02:
        errors.append(f"Unexpected slide size: {width:.3f} x {height:.3f}")

    all_text = []
    out_of_bounds = []
    chart_count = 0
    for slide_index, slide in enumerate(deck.slides, start=1):
        texts = []
        for shape_index, shape in enumerate(slide.shapes, start=1):
            text = shape_text(shape)
            if text:
                texts.append(text)
                all_text.append(text)
            left = shape.left / EMU_PER_INCH
            top = shape.top / EMU_PER_INCH
            right = (shape.left + shape.width) / EMU_PER_INCH
            bottom = (shape.top + shape.height) / EMU_PER_INCH
            intentional_decorative_bleed = not text and top >= -0.65 and right <= width + 0.6 and bottom <= height + 0.03
            if not intentional_decorative_bleed and (left < -0.03 or top < -0.03 or right > width + 0.03 or bottom > height + 0.03):
                out_of_bounds.append({
                    "slide": slide_index,
                    "shape": shape_index,
                    "box": [round(left, 3), round(top, 3), round(right, 3), round(bottom, 3)],
                    "text": text[:80],
                })
            if getattr(shape, "has_chart", False):
                chart_count += 1
        if slide_index not in (1, 34) and not texts:
            errors.append(f"Slide {slide_index} has no extracted text")

    if out_of_bounds:
        errors.append(f"Out-of-bounds shapes: {len(out_of_bounds)}")
    if chart_count < 5:
        errors.append(f"Expected at least 5 native charts, found {chart_count}")

    joined = "\n".join(all_text)
    required_phrases = [
        "超過100家店鋪表達使用意向",
        "尚無真實歸因GMV",
        "融資 2億日圓",
        "出讓 10% 股權",
        "一般方案",
        "激進方案",
        "保守情境",
        "透明點單",
        "CPS",
    ]
    for phrase in required_phrases:
        if phrase not in joined:
            errors.append(f"Missing required disclosure or phrase: {phrase}")
    forbidden_phrases = ["使用意向的店鋪已簽約", "100家已正式簽約", "保證回報", "成人性服務收入已納入"]
    for phrase in forbidden_phrases:
        if phrase in joined:
            errors.append(f"Forbidden overclaim found: {phrase}")

    if "时" in joined:
        errors.append("Simplified Chinese character found: 时")
    if "TODO" in joined or "FIXME" in joined:
        errors.append("Placeholder text found")

    report = {
        "file": str(args.pptx),
        "slides": len(deck.slides),
        "sizeInches": [round(width, 3), round(height, 3)],
        "nativeCharts": chart_count,
        "outOfBounds": out_of_bounds,
        "errors": errors,
        "warnings": warnings,
        "status": "pass" if not errors else "fail",
    }
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(0 if not errors else 1)


if __name__ == "__main__":
    main()
