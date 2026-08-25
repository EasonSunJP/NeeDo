#!/usr/bin/env python3
"""Structural and source verification for the NeeDo investor roadshow deck."""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader


EMU_PER_INCH = 914400
EXPECTED_CHART_SLIDES = [23, 24, 25, 27, 29]
REQUIRED_TEXT = [
    "100+ 店鋪為使用意向",
    "透明點單與 CPS 目前為可操作原型",
    "成人性服務相關業態現行條款禁止",
    "財務與投資回報為情境分析",
    "¥9,800",
    "約 1.53 個月",
    "約 3.06 個月",
    "一般方案",
    "激進方案",
    "¥200M",
    "10%",
    "投前估值",
    "投後估值",
]
FORBIDDEN_TEXT = [
    "SaaS月费0",
    "月額0",
    "Free｜月額0",
    "成人向業態可立即商用",
    "100+ 已簽約",
    "100+ 付費店鋪",
]


def shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False):
        return shape.text or ""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "\n".join(shape_text(child) for child in shape.shapes)
    return ""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--asset-manifest",
        type=Path,
        default=Path("assets/needo-roadshow-reference-style/manifest.json"),
    )
    args = parser.parse_args()

    errors: list[str] = []
    warnings: list[str] = []

    if not args.pptx.is_file():
        raise SystemExit(f"Missing PPTX: {args.pptx}")
    if not args.pdf.is_file():
        raise SystemExit(f"Missing PDF: {args.pdf}")

    with zipfile.ZipFile(args.pptx) as archive:
        bad_member = archive.testzip()
        if bad_member:
            errors.append(f"Corrupt ZIP member: {bad_member}")
        packed_names = "\n".join(archive.namelist()).lower()
        if "mascot" in packed_names:
            errors.append("Packed PPTX contains a mascot-named file")

    presentation = Presentation(args.pptx)
    slide_count = len(presentation.slides)
    width_in = presentation.slide_width / EMU_PER_INCH
    height_in = presentation.slide_height / EMU_PER_INCH
    if slide_count != 34:
        errors.append(f"Expected 34 slides, found {slide_count}")
    if abs(width_in / height_in - 16 / 9) > 0.002:
        errors.append(f"Expected 16:9, found {width_in:.3f} x {height_in:.3f}")

    chart_slides: list[int] = []
    all_text_parts: list[str] = []
    out_of_bounds: list[dict[str, object]] = []
    tolerance = 0.02
    for slide_index, slide in enumerate(presentation.slides, start=1):
        if any(shape.shape_type == MSO_SHAPE_TYPE.CHART for shape in slide.shapes):
            chart_slides.append(slide_index)
        for shape in slide.shapes:
            all_text_parts.append(shape_text(shape))
            x = shape.left / EMU_PER_INCH
            y = shape.top / EMU_PER_INCH
            w = shape.width / EMU_PER_INCH
            h = shape.height / EMU_PER_INCH
            if x < -tolerance or y < -tolerance or x + w > width_in + tolerance or y + h > height_in + tolerance:
                out_of_bounds.append({
                    "slide": slide_index,
                    "name": shape.name,
                    "bounds": [round(x, 3), round(y, 3), round(w, 3), round(h, 3)],
                })

    all_text = "\n".join(all_text_parts)
    for required in REQUIRED_TEXT:
        if required not in all_text:
            errors.append(f"Missing required text: {required}")
    for forbidden in FORBIDDEN_TEXT:
        if forbidden in all_text:
            errors.append(f"Forbidden legacy claim found: {forbidden}")
    if chart_slides != EXPECTED_CHART_SLIDES:
        errors.append(f"Native charts expected on {EXPECTED_CHART_SLIDES}, found {chart_slides}")
    if out_of_bounds:
        errors.append(f"Found {len(out_of_bounds)} out-of-bounds top-level shapes")

    pdf_pages = len(PdfReader(str(args.pdf)).pages)
    if pdf_pages != 34:
        errors.append(f"Expected 34 PDF pages, found {pdf_pages}")

    manifest = json.loads(args.asset_manifest.read_text(encoding="utf-8"))
    asset_dir = args.asset_manifest.parent
    asset_checks: list[dict[str, object]] = []
    for item in manifest:
        asset_path = asset_dir / item["file"]
        matches = asset_path.is_file() and sha256(asset_path) == item["sha256"]
        mascot_removed = item.get("mascotRemoved") is True
        asset_checks.append({
            "id": item["id"],
            "file": str(asset_path),
            "hashMatches": matches,
            "mascotRemoved": mascot_removed,
        })
        if not matches:
            errors.append(f"Asset hash mismatch: {item['file']}")
        if not mascot_removed:
            errors.append(f"Mascot removal not attested: {item['file']}")

    report = {
        "status": "passed" if not errors else "failed",
        "pptx": str(args.pptx.resolve()),
        "pdf": str(args.pdf.resolve()),
        "slides": slide_count,
        "pdfPages": pdf_pages,
        "sizeInches": [round(width_in, 3), round(height_in, 3)],
        "nativeChartSlides": chart_slides,
        "outOfBounds": out_of_bounds,
        "assets": asset_checks,
        "warnings": warnings,
        "errors": errors,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
