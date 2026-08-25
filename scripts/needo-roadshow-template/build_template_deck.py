#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = Path("/Users/eason/Documents/NeeDo/PPT模板/PPT模板.pptx")
DEFAULT_OUTPUT = ROOT / "outputs/needo-roadshow-template-2026-08-23/NeeDo_海外投資人路演_BP_繁中_模板版_2026-08-23.pptx"
EXPORTER = ROOT / "scripts/needo-roadshow-template/export-data.mjs"

W = 13.333
H = 7.5
FONT = "Arial Unicode MS"

COLORS = {
    "dark": "455F51",
    "green": "549E39",
    "lime": "8AB833",
    "yellow": "C0CF3A",
    "teal": "029676",
    "cyan": "4AB5C4",
    "blue": "0989B1",
    "beige": "E3DED1",
    "light": "F5F7F1",
    "soft": "EEF3E7",
    "ink": "242B27",
    "muted": "69736D",
    "line": "DDE4D9",
    "white": "FFFFFF",
    "risk": "A56748",
}

TEMPLATE_PAGE_MAP = [
    1, 33, 2, 34, 31, 11, 8, 32, 36,
    12, 14, 20, 23, 38, 24, 16, 35, 40,
    7, 26, 19, 17, 3, 4, 13, 10, 25, 37,
    29, 30, 22, 39, 6, 42,
]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def inch(value: float):
    return Inches(value)


def load_payload() -> dict:
    result = subprocess.run(
        ["node", str(EXPORTER)],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(result.stdout)


def set_slide_order(prs: Presentation, source_pages: list[int]) -> None:
    if len(source_pages) != len(set(source_pages)):
        raise ValueError("Template mapping must use unique source slides")
    sld_id_list = prs.slides._sldIdLst
    original = list(sld_id_list)
    selected = [original[index - 1] for index in source_pages]
    selected_ids = {id(item) for item in selected}
    for item in original:
        if id(item) not in selected_ids:
            prs.part.drop_rel(item.rId)
    for item in original:
        sld_id_list.remove(item)
    for item in selected:
        sld_id_list.append(item)


def remove_all_shapes(slide) -> None:
    for shape in list(slide.shapes):
        element = shape._element
        element.getparent().remove(element)


def remove_negative_position_shapes(slide) -> None:
    for shape in list(slide.shapes):
        if shape.left < 0 or shape.top < 0:
            element = shape._element
            element.getparent().remove(element)


def clear_text_recursive(shapes) -> None:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            clear_text_recursive(shape.shapes)
        elif getattr(shape, "has_text_frame", False):
            shape.text_frame.clear()


def set_shape_fill(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)


def set_shape_line(shape, color: str, width: float = 1.0) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def add_shape(slide, shape_type, x, y, w, h, fill="FFFFFF", line="DDE4D9", radius=False):
    shape = slide.shapes.add_shape(shape_type, inch(x), inch(y), inch(w), inch(h))
    set_shape_fill(shape, fill)
    set_shape_line(shape, line, 0.8)
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=12,
    color="242B27",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    font=FONT,
    margin=0.03,
    fit=True,
):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = inch(margin)
    frame.margin_right = inch(margin)
    frame.margin_top = inch(margin)
    frame.margin_bottom = inch(margin)
    if fit:
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=10, color="242B27", gap=1.0):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = inch(0.04)
    frame.margin_top = frame.margin_bottom = inch(0.03)
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(gap)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
    return box


def add_pill(slide, text, x, y, w, h=0.36, fill="549E39", color="FFFFFF", size=9.5):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, fill)
    add_text(slide, text, x + 0.04, y + 0.02, w - 0.08, h - 0.04, size, color, True, PP_ALIGN.CENTER)


def add_card(slide, x, y, w, h, fill="FFFFFF", line="DDE4D9"):
    return add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line)


def add_circle(slide, x, y, d, fill, line=None):
    return add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, d, d, fill, line or fill)


def add_ring(slide, x, y, d, outer, inner="FFFFFF", thickness=0.18):
    add_circle(slide, x, y, d, outer)
    add_circle(slide, x + thickness, y + thickness, d - 2 * thickness, inner, inner)


def begin_slide(slide, page: int, meta: dict, kicker="NeeDo｜海外投資人路演", page_label=None):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(COLORS["white"])
    add_text(slide, kicker, 0.55, 0.20, 3.0, 0.22, 8.5, COLORS["green"], True)
    add_text(slide, meta["title"], 0.62, 0.44, 12.1, 0.52, 19, COLORS["ink"], True, PP_ALIGN.CENTER)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.62, 1.02, 12.1, 0.025, COLORS["line"], COLORS["line"])
    page_text = f"{page:02d}" if page_label is None else page_label
    if page_text:
        add_text(slide, page_text, 12.25, 0.18, 0.5, 0.25, 8.5, COLORS["green"], True, PP_ALIGN.RIGHT)


def add_disclosure(slide, meta: dict, source_prefix="來源"):
    text = meta["statement"]
    size = 7.2 if len(text) < 180 else 6.5
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.65, 6.52, 12.05, 0.47, COLORS["light"], COLORS["light"])
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.65, 6.52, 0.06, 0.47, COLORS["green"], COLORS["green"])
    add_text(slide, text, 0.80, 6.57, 11.70, 0.33, size, COLORS["muted"], False)
    add_text(slide, f"{source_prefix}：{meta['source']}", 0.68, 7.10, 8.5, 0.18, 6.8, COLORS["muted"])
    add_text(slide, "NeeDo｜機密・投資人討論用", 9.4, 7.10, 3.25, 0.18, 6.8, COLORS["muted"], False, PP_ALIGN.RIGHT)


def add_metric(slide, value, label, x, y, w, accent="549E39", note=None):
    add_card(slide, x, y, w, 1.14, COLORS["white"], COLORS["line"])
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, 0.08, 1.14, accent, accent)
    add_text(slide, value, x + 0.20, y + 0.15, w - 0.35, 0.42, 24, accent, True)
    add_text(slide, label, x + 0.20, y + 0.61, w - 0.35, 0.23, 9.2, COLORS["ink"], True)
    if note:
        add_text(slide, note, x + 0.20, y + 0.84, w - 0.35, 0.18, 7.2, COLORS["muted"])


def style_chart(chart, colors, show_legend=True, labels=True):
    chart.has_legend = show_legend
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = FONT
        chart.legend.font.size = Pt(8)
    chart.chart_style = 10
    try:
        chart.category_axis.tick_labels.font.name = FONT
        chart.category_axis.tick_labels.font.size = Pt(7)
        chart.value_axis.tick_labels.font.name = FONT
        chart.value_axis.tick_labels.font.size = Pt(8)
        chart.value_axis.has_major_gridlines = True
    except ValueError:
        pass
    for index, series in enumerate(chart.series):
        color = colors[index % len(colors)]
        try:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = rgb(color)
        except Exception:
            series.format.line.color.rgb = rgb(color)
        if labels:
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.font.name = FONT
            plot.data_labels.font.size = Pt(7)
            plot.data_labels.font.color.rgb = rgb(COLORS["muted"])
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.number_format = "0.0;-0.0"


def add_column_chart(slide, categories, series, x, y, w, h, colors, minimum=None, maximum=None, major=None, labels=True):
    data = ChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        inch(x), inch(y), inch(w), inch(h), data,
    ).chart
    style_chart(chart, colors, True, labels)
    if minimum is not None:
        chart.value_axis.minimum_scale = minimum
    if maximum is not None:
        chart.value_axis.maximum_scale = maximum
    if major is not None:
        chart.value_axis.major_unit = major
    return chart


def add_line_chart(slide, categories, series, x, y, w, h, colors):
    data = ChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        inch(x), inch(y), inch(w), inch(h), data,
    ).chart
    style_chart(chart, colors, True, False)
    for index, series_item in enumerate(chart.series):
        series_item.format.line.color.rgb = rgb(colors[index % len(colors)])
        series_item.format.line.width = Pt(2.25)
        series_item.marker.style = XL_MARKER_STYLE.CIRCLE
        series_item.marker.size = 5
    return chart


def add_doughnut(slide, labels, values, x, y, w, h, colors):
    data = ChartData()
    data.categories = labels
    data.add_series("資金配置", values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT,
        inch(x), inch(y), inch(w), inch(h), data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.font.name = FONT
    chart.legend.font.size = Pt(7.5)
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.show_percentage = True
    chart.plots[0].data_labels.show_value = False
    chart.plots[0].data_labels.font.name = FONT
    chart.plots[0].data_labels.font.size = Pt(8)
    chart.plots[0].data_labels.font.color.rgb = rgb(COLORS["dark"])
    chart.doughnut_hole_size = 62
    for index, point in enumerate(chart.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(colors[index % len(colors)])
    return chart


def fmt_int(value):
    return f"{int(value):,}"


def fmt_m(value):
    return f"¥{value:,.1f}M"


def page_1(slide, meta, data):
    remove_negative_position_shapes(slide)
    clear_text_recursive(slide.shapes)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 3.15, 1.15, 5.25, 2.45, COLORS["white"], COLORS["white"])
    add_text(slide, "NeeDo", 3.42, 1.34, 4.70, 0.55, 31, COLORS["green"], True, PP_ALIGN.CENTER)
    add_text(slide, "日本生活服務履約網路", 3.42, 1.94, 4.70, 0.40, 18, COLORS["ink"], True, PP_ALIGN.CENTER)
    add_text(slide, "海外投資人路演｜繁體中文母版", 3.42, 2.42, 4.70, 0.30, 11, COLORS["muted"], False, PP_ALIGN.CENTER)
    add_text(slide, "Pre-A｜融資 2億日圓｜出讓 10% 股權", 3.42, 2.92, 4.70, 0.30, 12.5, COLORS["green"], True, PP_ALIGN.CENTER)
    add_text(slide, "讓服務被看見、被選擇、被完成、被分配", 3.35, 4.72, 5.15, 0.40, 14.5, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_text(slide, "2026年8月｜投資人討論用", 4.15, 5.22, 3.55, 0.23, 8, COLORS["green"], False, PP_ALIGN.CENTER)


def page_2(slide, meta, data):
    begin_slide(slide, 2, meta)
    add_text(slide, "NeeDo 將分散、人工、不可計算的服務業交易，轉化為可追蹤、可調度、可歸因的履約網路。", 1.1, 1.28, 11.1, 0.72, 19, COLORS["dark"], True, PP_ALIGN.CENTER)
    judgments = [
        ("01", "市場拉力", "訪日與在留外國人持續增長"),
        ("02", "結構瓶頸", "五方即時協調仍靠電話與LINE"),
        ("03", "產品解法", "AI排班＋透明點單＋CPS事件鏈"),
        ("04", "資本槓桿", "2億日圓換取付費、留存與城市密度驗證"),
    ]
    for idx, (n, title, body) in enumerate(judgments):
        x = 0.90 + idx * 3.05
        add_pill(slide, n, x, 2.40, 0.60, 0.38, [COLORS["green"], COLORS["lime"], COLORS["yellow"], COLORS["teal"]][idx])
        add_text(slide, title, x, 2.95, 2.55, 0.32, 13, COLORS["ink"], True, PP_ALIGN.CENTER)
        add_text(slide, body, x, 3.40, 2.55, 0.75, 9.5, COLORS["muted"], False, PP_ALIGN.CENTER)
        if idx < 3:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, x + 2.60, 3.28, 0.38, 0.50, COLORS["beige"], COLORS["beige"])
    add_metric(slide, "¥200M", "本輪融資", 1.20, 4.75, 2.45, COLORS["green"], "Pre-A")
    add_metric(slide, "10%", "出讓股權", 3.85, 4.75, 2.45, COLORS["lime"], "交割後口徑")
    add_metric(slide, "¥1.8B", "投前估值", 6.50, 4.75, 2.45, COLORS["yellow"], "Pre-money")
    add_metric(slide, "¥2.0B", "投後估值", 9.15, 4.75, 2.45, COLORS["teal"], "Post-money")
    add_disclosure(slide, meta)


def page_3(slide, meta, data):
    begin_slide(slide, 3, meta)
    market = data["market"]
    add_metric(slide, f"{market['visitors2025']/10_000:,.2f}萬", "2025年訪日外客", 1.0, 1.35, 5.25, COLORS["green"], f"年增 +{market['visitorsGrowth']}%｜S1")
    add_metric(slide, f"{market['foreignResidents2025']/10_000:,.2f}萬", "2025年末在留外國人", 7.05, 1.35, 5.25, COLORS["teal"], f"年增 +{market['foreignResidentsGrowth']}%｜S2")
    add_text(slide, "需求不是單一旅遊旺季，而是『訪日＋在留』兩條長期增量曲線。", 2.05, 2.92, 9.25, 0.48, 16, COLORS["dark"], True, PP_ALIGN.CENTER)
    for idx, (label, value, color) in enumerate([
        ("跨語言預約", "入口摩擦", COLORS["green"]),
        ("可履約供給", "協調摩擦", COLORS["lime"]),
        ("價格與規則", "信任摩擦", COLORS["yellow"]),
        ("變更與取消", "保障摩擦", COLORS["teal"]),
    ]):
        x = 1.10 + idx * 3.0
        add_circle(slide, x, 4.02, 0.66, color)
        add_text(slide, str(idx + 1), x, 4.15, 0.66, 0.25, 11, COLORS["white"], True, PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.82, 4.00, 1.85, 0.28, 11, COLORS["ink"], True)
        add_text(slide, value, x + 0.82, 4.35, 1.85, 0.22, 8.5, COLORS["muted"])
    add_text(slide, "投資含義：多語言入口只是第一步，真正價值在把需求轉成可完成的服務交易。", 1.25, 5.35, 10.85, 0.50, 14, COLORS["green"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_4(slide, meta, data):
    begin_slide(slide, 4, meta)
    steps = ["客人電話詢問", "店鋪確認需求", "LINE／電話聯絡技師", "技師回覆行程", "店鋪回覆客人", "說明規則與場所"]
    colors = [COLORS["green"], COLORS["green"], COLORS["lime"], COLORS["yellow"], COLORS["teal"], COLORS["blue"]]
    for idx, step in enumerate(steps):
        x = 0.75 + idx * 2.05
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, x, 1.65, 1.85, 0.75, colors[idx], colors[idx])
        add_text(slide, step, x + 0.24, 1.82, 1.32, 0.35, 8.2, COLORS["white"], True, PP_ALIGN.CENTER)
    add_text(slide, "變更一次＝流程重做一次", 2.35, 3.18, 8.65, 0.70, 25, COLORS["dark"], True, PP_ALIGN.CENTER)
    frictions = ["語言限制", "多次人工轉述", "行程變更成本倍增", "場所與規則說明繁瑣", "取消與爭議難追蹤"]
    for idx, item in enumerate(frictions):
        add_pill(slide, item, 0.95 + idx * 2.45, 4.32, 2.08, 0.50, COLORS["light"], COLORS["dark"], 9)
    add_text(slide, "結論：預約形式多年未見本質改變，電話與LINE仍是營運主幹。", 1.20, 5.35, 10.95, 0.45, 14, COLORS["green"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_5(slide, meta, data):
    begin_slide(slide, 5, meta)
    center_x, center_y, center_d = 5.55, 2.25, 2.15
    add_circle(slide, center_x, center_y, center_d, COLORS["green"])
    add_text(slide, "即時協調\n＋\n消費透明", center_x + 0.20, center_y + 0.48, center_d - 0.40, 1.15, 15, COLORS["white"], True, PP_ALIGN.CENTER)
    nodes = [
        ("店鋪", 1.20, 1.55, COLORS["green"]),
        ("中介", 9.90, 1.45, COLORS["lime"]),
        ("服務者", 10.15, 4.25, COLORS["yellow"]),
        ("客人", 1.15, 4.25, COLORS["teal"]),
        ("場所", 5.65, 4.60, COLORS["blue"]),
    ]
    for label, x, y, color in nodes:
        add_ring(slide, x, y, 1.18, color, COLORS["white"], 0.15)
        add_text(slide, label, x + 0.10, y + 0.42, 0.98, 0.28, 11, color, True, PP_ALIGN.CENTER)
        line = slide.shapes.add_connector(1, inch(x + 0.59), inch(y + 0.59), inch(center_x + center_d/2), inch(center_y + center_d/2))
        set_shape_line(line, COLORS["line"], 1.2)
    add_text(slide, "高消費場所可透過 NeeDo 點單系統，先選擇、再確認、即時看總額，降低不透明消費風險。", 2.05, 5.92, 9.25, 0.35, 10.5, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_6(slide, meta, data):
    begin_slide(slide, 6, meta)
    roles = [
        ("店鋪", "找人、確認、再通知", COLORS["green"]),
        ("中介", "逐一聯絡客人與技師", COLORS["lime"]),
        ("服務者", "回覆行程、變更與場所", COLORS["yellow"]),
        ("客人", "重複說明、等待與承擔不確定", COLORS["teal"]),
    ]
    for idx, (role, body, color) in enumerate(roles):
        x = 0.95 + idx * 3.08
        add_ring(slide, x + 0.60, 1.55, 1.32, color, COLORS["white"], 0.18)
        add_text(slide, role, x + 0.72, 2.00, 1.08, 0.30, 12, color, True, PP_ALIGN.CENTER)
        add_text(slide, body, x, 3.15, 2.55, 0.66, 10, COLORS["ink"], True, PP_ALIGN.CENTER)
        add_pill(slide, "行程一變，人工成本翻倍", x + 0.15, 4.14, 2.25, 0.46, COLORS["light"], COLORS["dark"], 8.4)
    add_text(slide, "平台外私加聯絡方式，會造成跳單、隱私外洩與責任邊界消失。", 1.35, 5.20, 10.65, 0.50, 15, COLORS["green"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_7(slide, meta, data):
    begin_slide(slide, 7, meta)
    stages = [
        ("有需求", 1.00, 0.95, COLORS["green"]),
        ("日語限定預約", 3.20, 0.78, COLORS["lime"]),
        ("拒絕外國客", 5.40, 0.61, COLORS["yellow"]),
        ("規則／價格不清", 7.60, 0.45, COLORS["teal"]),
        ("取消或求償困難", 9.80, 0.30, COLORS["blue"]),
    ]
    for idx, (label, x, ratio, color) in enumerate(stages):
        w = 1.50 + ratio * 2.0
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 1.52 + idx * 0.78, w, 0.52, color, color)
        add_text(slide, label, x + 0.10, 1.63 + idx * 0.78, w - 0.20, 0.28, 9.8, COLORS["white"], True)
        if idx > 0:
            add_text(slide, "退出點", x - 0.78, 1.66 + idx * 0.78, 0.62, 0.20, 7.5, COLORS["risk"], True, PP_ALIGN.RIGHT)
    add_card(slide, 8.95, 2.00, 3.10, 2.78, COLORS["light"], COLORS["line"])
    add_text(slide, "結構性觀察", 9.30, 2.30, 2.40, 0.30, 13, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_rich_lines(slide, ["• 並非所有店鋪都拒絕外國客", "• 問題集中在語言、責任與營運成本", "• 多語言入口必須連接可履約供給", "• 透明規則與爭議處理同樣重要"], 9.25, 2.86, 2.50, 1.55, 8.8, COLORS["ink"], 3)
    add_disclosure(slide, meta)


def page_8(slide, meta, data):
    begin_slide(slide, 8, meta)
    tiers = [
        ("生活服務基礎設施", 1.40, 1.40, 10.50, 1.20, COLORS["green"]),
        ("相鄰預約／派遣服務", 2.45, 2.65, 8.40, 1.20, COLORS["lime"]),
        ("高摩擦首發市場", 3.55, 3.90, 6.20, 1.20, COLORS["yellow"]),
    ]
    for label, x, y, w, h, color in tiers:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.TRAPEZOID, x, y, w, h, color, color)
        add_text(slide, label, x + 0.40, y + 0.37, w - 0.80, 0.34, 13, COLORS["white"], True, PP_ALIGN.CENTER)
    market = data["market"]
    metrics = [
        (fmt_int(market["receptionVenues"]), "接待飲食營業據點"),
        (fmt_int(market["lateNightAlcoholVenues"]), "深夜酒類提供據點"),
        (fmt_int(market["dispatchAdultBusinesses"]), "派遣型成人特殊營業脈絡"),
    ]
    for idx, (value, label) in enumerate(metrics):
        add_metric(slide, value, label, 0.95 + idx * 4.05, 5.25, 3.48, [COLORS["green"], COLORS["lime"], COLORS["teal"]][idx], "官方供給側脈絡；不等於核心TAM")
    add_disclosure(slide, meta)


def page_9(slide, meta, data):
    begin_slide(slide, 9, meta)
    zones = [
        ("可立即進入", "一般生活服務、店鋪SaaS、預約與排班", "允許", COLORS["green"]),
        ("完成控制後進入", "按摩資格、場所規則、支付／退款與責任邊界", "條件式", COLORS["yellow"]),
        ("現行禁止", "成人性服務；未來須獨立產品、法務與監管評估", "隔離", COLORS["teal"]),
    ]
    for idx, (title, body, status, color) in enumerate(zones):
        y = 1.45 + idx * 1.48
        add_card(slide, 1.15, y, 11.05, 1.02, COLORS["white"], COLORS["line"])
        add_pill(slide, status, 1.42, y + 0.24, 1.10, 0.44, color, COLORS["white"], 9.5)
        add_text(slide, title, 2.85, y + 0.16, 2.40, 0.30, 13, color, True)
        add_text(slide, body, 5.05, y + 0.14, 6.55, 0.50, 9.5, COLORS["ink"], True)
    add_text(slide, "合規不是附錄，而是產品的進入順序。", 2.00, 5.75, 9.35, 0.42, 17, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_10(slide, meta, data):
    begin_slide(slide, 10, meta)
    add_circle(slide, 5.23, 2.20, 2.90, COLORS["green"])
    add_text(slide, "NeeDo\n履約核心", 5.60, 3.00, 2.15, 0.72, 19, COLORS["white"], True, PP_ALIGN.CENTER)
    modules = [
        ("多語言預約", 1.10, 1.65, COLORS["green"]),
        ("AI排班", 9.70, 1.65, COLORS["lime"]),
        ("透明點單", 1.10, 4.58, COLORS["yellow"]),
        ("CPS歸因", 9.70, 4.58, COLORS["teal"]),
    ]
    for label, x, y, color in modules:
        add_ring(slide, x, y, 1.55, color, COLORS["white"], 0.20)
        add_text(slide, label, x + 0.18, y + 0.57, 1.18, 0.35, 11, color, True, PP_ALIGN.CENTER)
        line = slide.shapes.add_connector(1, inch(x + 0.78), inch(y + 0.78), inch(6.68), inch(3.65))
        set_shape_line(line, COLORS["line"], 1.2)
    add_pill(slide, "受保護溝通與權益", 4.80, 5.48, 3.75, 0.50, COLORS["dark"], COLORS["white"], 11)
    add_disclosure(slide, meta)


def page_11(slide, meta, data):
    begin_slide(slide, 11, meta)
    center = (6.67, 3.50)
    add_ring(slide, 4.90, 1.72, 3.54, COLORS["green"], COLORS["white"], 0.28)
    add_text(slide, "可追蹤\n履約閉環", 5.68, 2.95, 1.98, 0.85, 18, COLORS["dark"], True, PP_ALIGN.CENTER)
    stages = ["內容／發現", "選擇", "預約／確認", "調整／履約", "保障／支付", "歸因／結算"]
    for idx, label in enumerate(stages):
        angle = -math.pi / 2 + idx * math.pi / 3
        x = center[0] + math.cos(angle) * 4.28 - 0.65
        y = center[1] + math.sin(angle) * 2.25 - 0.32
        color = [COLORS["green"], COLORS["lime"], COLORS["yellow"], COLORS["teal"], COLORS["cyan"], COLORS["blue"]][idx]
        add_pill(slide, f"{idx+1}. {label}", x, y, 1.55, 0.56, color, COLORS["white"], 8.8)
    add_disclosure(slide, meta)


def page_12(slide, meta, data):
    begin_slide(slide, 12, meta)
    rows = ["技能匹配", "可用時段", "場所規則", "移動時間", "衝突檢查", "替補優先序"]
    colors = [COLORS["green"], COLORS["green"], COLORS["lime"], COLORS["yellow"], COLORS["teal"], COLORS["blue"]]
    add_text(slide, "09:00", 3.30, 1.25, 1.0, 0.25, 8, COLORS["muted"], False, PP_ALIGN.CENTER)
    add_text(slide, "12:00", 6.20, 1.25, 1.0, 0.25, 8, COLORS["muted"], False, PP_ALIGN.CENTER)
    add_text(slide, "18:00", 9.25, 1.25, 1.0, 0.25, 8, COLORS["muted"], False, PP_ALIGN.CENTER)
    for idx, row in enumerate(rows):
        y = 1.70 + idx * 0.68
        add_text(slide, row, 0.90, y + 0.08, 1.75, 0.25, 10, COLORS["ink"], True, PP_ALIGN.RIGHT)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 2.85, y, 8.95, 0.42, COLORS["beige"], COLORS["beige"])
        start = 3.0 + (idx % 3) * 0.78
        width = 3.75 + ((idx + 1) % 3) * 0.92
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, start, y, width, 0.42, colors[idx], colors[idx])
    outcomes = ["變更可重算", "候補可排序", "跨供給可調度"]
    for idx, item in enumerate(outcomes):
        add_pill(slide, item, 2.05 + idx * 3.45, 5.85, 2.85, 0.45, [COLORS["green"], COLORS["lime"], COLORS["teal"]][idx], COLORS["white"], 10)
    add_disclosure(slide, meta)


def page_13(slide, meta, data):
    begin_slide(slide, 13, meta)
    layers = [
        ("店內服務者", "優先滿足；保留既有關係與責任", 9.60, COLORS["green"]),
        ("合作店供給", "跨店共享；需權限、距離與店規匹配", 7.45, COLORS["lime"]),
        ("外部服務者", "只在資格、責任與合規審核清楚時啟用", 5.30, COLORS["teal"]),
    ]
    for idx, (label, body, width, color) in enumerate(layers):
        y = 1.55 + idx * 1.28
        add_text(slide, label, 0.95, y + 0.16, 1.70, 0.28, 11, color, True, PP_ALIGN.RIGHT)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 2.95, y, width, 0.66, color, color)
        add_text(slide, body, 3.20, y + 0.16, width - 0.50, 0.28, 9.5, COLORS["white"], True)
    add_text(slide, "AI 產生建議，不取代店鋪最終授權；例外與變更保留事件紀錄。", 1.45, 5.45, 10.45, 0.50, 15, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_14(slide, meta, data):
    begin_slide(slide, 14, meta)
    add_text(slide, "平台外", 1.20, 1.30, 4.80, 0.36, 15, COLORS["risk"], True, PP_ALIGN.CENTER)
    add_text(slide, "平台內", 7.30, 1.30, 4.80, 0.36, 15, COLORS["green"], True, PP_ALIGN.CENTER)
    left = ["私加聯絡方式", "跳過店鋪／中介", "隱私與身份外洩", "變更與爭議無紀錄"]
    right = ["角色最小化資訊", "平台訊息與訂單事件", "權益／激勵降低跳單", "可追蹤變更與爭議處理"]
    for idx in range(4):
        y = 1.92 + idx * 0.88
        add_pill(slide, left[idx], 1.20, y, 4.80, 0.55, COLORS["light"], COLORS["risk"], 10)
        add_pill(slide, right[idx], 7.30, y, 4.80, 0.55, COLORS["soft"], COLORS["green"], 10)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, 6.25, y + 0.05, 0.72, 0.46, COLORS["beige"], COLORS["beige"])
    add_text(slide, "保護隱私與分配權，不靠封鎖聯絡，而靠可追蹤的權益設計。", 1.50, 5.68, 10.35, 0.42, 15, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_15(slide, meta, data):
    begin_slide(slide, 15, meta)
    add_card(slide, 1.00, 1.32, 7.25, 4.75, COLORS["white"], COLORS["line"])
    add_text(slide, "透明點單示意", 1.35, 1.62, 3.00, 0.30, 14, COLORS["green"], True)
    items = [
        ("基礎服務", "已選擇", "¥12,000", COLORS["green"]),
        ("追加項目 A", "待確認", "¥3,000", COLORS["yellow"]),
        ("追加項目 B", "已確認", "¥2,000", COLORS["teal"]),
        ("場所費／其他", "明細顯示", "¥0", COLORS["blue"]),
    ]
    for idx, (name, status, price, color) in enumerate(items):
        y = 2.15 + idx * 0.70
        add_text(slide, name, 1.42, y, 2.20, 0.28, 10, COLORS["ink"], True)
        add_pill(slide, status, 3.80, y - 0.03, 1.10, 0.34, color, COLORS["white"], 7.8)
        add_text(slide, price, 6.35, y, 1.40, 0.28, 11, color, True, PP_ALIGN.RIGHT)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 1.40, 5.08, 6.35, 0.025, COLORS["line"], COLORS["line"])
    add_text(slide, "即時累計", 1.42, 5.30, 2.20, 0.30, 11, COLORS["muted"], True)
    add_text(slide, "¥17,000", 5.50, 5.18, 2.20, 0.50, 23, COLORS["green"], True, PP_ALIGN.RIGHT)
    add_card(slide, 8.72, 1.32, 3.55, 4.75, COLORS["light"], COLORS["line"])
    add_text(slide, "先選、再確認、全程看總額", 9.08, 1.78, 2.82, 0.78, 18, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_rich_lines(slide, ["✓ 含稅價格", "✓ 可選項", "✓ 追加確認", "✓ 即時累計", "✓ 明細帳單"], 9.20, 2.85, 2.60, 1.90, 11, COLORS["green"], 4)
    add_text(slide, "金額僅為介面示意", 9.15, 5.40, 2.70, 0.24, 7.5, COLORS["muted"], False, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_16(slide, meta, data):
    begin_slide(slide, 16, meta)
    add_ring(slide, 4.50, 1.38, 4.35, COLORS["green"], COLORS["white"], 0.34)
    add_ring(slide, 5.12, 2.00, 3.10, COLORS["teal"], COLORS["white"], 0.30)
    add_text(slide, "NeeDo", 5.72, 3.10, 1.90, 0.45, 19, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_pill(slide, "AI 排班引擎", 1.10, 2.08, 2.55, 0.62, COLORS["green"], COLORS["white"], 12)
    add_pill(slide, "CPS 歸因引擎", 9.68, 2.08, 2.55, 0.62, COLORS["teal"], COLORS["white"], 12)
    add_text(slide, "把供給變成可履約容量", 1.20, 2.92, 2.30, 0.42, 9.5, COLORS["muted"], True, PP_ALIGN.CENTER)
    add_text(slide, "把內容變成可歸因交易", 9.80, 2.92, 2.30, 0.42, 9.5, COLORS["muted"], True, PP_ALIGN.CENTER)
    add_text(slide, "服務 × 達人 × 預約 × 履約 × 結算", 3.15, 5.55, 7.05, 0.50, 17, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_17(slide, meta, data):
    begin_slide(slide, 17, meta)
    steps = [
        ("內容", COLORS["green"]), ("達人", COLORS["green"]), ("預約", COLORS["lime"]),
        ("履約完成", COLORS["yellow"]), ("退款／風控", COLORS["risk"]), ("合格佣金結算", COLORS["teal"]),
    ]
    for idx, (label, color) in enumerate(steps):
        y = 1.42 + idx * 0.78
        add_pill(slide, f"{idx+1:02d}", 1.05, y, 0.70, 0.45, color, COLORS["white"], 9)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 2.05, y, 8.70, 0.45, COLORS["light"], COLORS["light"])
        add_text(slide, label, 2.35, y + 0.05, 3.10, 0.28, 10.5, color, True)
        add_text(slide, ["來源事件", "創作者識別", "服務與店鋪", "完成狀態", "取消／退款／資格", "只對合格交易結算"][idx], 6.10, y + 0.05, 4.20, 0.28, 8.8, COLORS["muted"], True)
        if idx < len(steps) - 1:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.DOWN_ARROW, 10.98, y + 0.48, 0.40, 0.26, COLORS["beige"], COLORS["beige"])
    add_disclosure(slide, meta)


def page_18(slide, meta, data):
    begin_slide(slide, 18, meta)
    levels = [
        ("已可操作", COLORS["green"], ["多端體驗原型", "預約／排班流程", "透明點單介面", "CPS事件設計"]),
        ("本輪完成", COLORS["lime"], ["正式後端／資料庫", "權限／支付／審計", "真實付費與留存", "跨店供給調度"]),
        ("未宣稱", COLORS["teal"], ["已規模商業化", "已產生CPS收入", "100+已簽約／付費", "估值或退出保證"]),
    ]
    for idx, (title, color, items) in enumerate(levels):
        x = 0.95 + idx * 4.10
        add_card(slide, x, 1.38, 3.55, 4.58, COLORS["white"], color)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, 1.38, 3.55, 0.72, color, color)
        add_text(slide, title, x + 0.15, 1.58, 3.25, 0.30, 15, COLORS["white"], True, PP_ALIGN.CENTER)
        for item_idx, item in enumerate(items):
            add_pill(slide, item, x + 0.34, 2.42 + item_idx * 0.73, 2.87, 0.47, COLORS["light"], COLORS["dark"], 8.6)
    add_disclosure(slide, meta)


def page_19(slide, meta, data):
    begin_slide(slide, 19, meta)
    add_text(slide, "100+", 1.00, 1.35, 5.40, 1.55, 70, COLORS["green"], True, PP_ALIGN.CENTER)
    add_text(slide, "店鋪使用意向", 1.05, 3.02, 5.35, 0.50, 22, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_card(slide, 7.00, 1.48, 5.10, 3.55, COLORS["light"], COLORS["line"])
    add_text(slide, "這個數字代表什麼", 7.42, 1.85, 4.30, 0.34, 14, COLORS["green"], True)
    add_rich_lines(slide, ["✓ 商談中明確表達使用意向", "✓ 反映痛點存在與產品方向吸引力", "✕ 不等於已簽約", "✕ 不等於已付費／活躍店鋪", "✕ 不等於收入或GMV"], 7.40, 2.38, 4.20, 1.95, 10.5, COLORS["ink"], 4)
    add_text(slide, "下一步證據：CRM → 試點 → 合約 → 收款 → 留存", 1.25, 5.45, 10.75, 0.52, 15, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_20(slide, meta, data):
    begin_slide(slide, 20, meta)
    stairs = [
        ("單城／單一高摩擦業態", 1.0, 4.65, 2.45, 0.72, COLORS["green"]),
        ("建立店鋪與服務者密度", 3.20, 3.75, 2.75, 1.62, COLORS["lime"]),
        ("驗證履約與留存", 5.70, 2.85, 2.75, 2.52, COLORS["yellow"]),
        ("複製到第二城市／業態", 8.20, 1.95, 3.55, 3.42, COLORS["teal"]),
    ]
    for idx, (label, x, y, w, h, color) in enumerate(stairs):
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h, color, color)
        add_text(slide, f"0{idx+1}\n{label}", x + 0.18, y + 0.18, w - 0.36, min(0.70, h - 0.25), 10, COLORS["white"], True, PP_ALIGN.CENTER)
    kpis = ["意向→上線轉換", "可服務時段覆蓋", "履約完成率", "跨店調度成功率"]
    for idx, item in enumerate(kpis):
        add_pill(slide, item, 1.05 + idx * 2.95, 5.75, 2.45, 0.44, COLORS["light"], COLORS["dark"], 8.6)
    add_text(slide, "以上為融資後逐城驗證指標，尚非已達成KPI。", 8.75, 5.80, 3.10, 0.28, 7.8, COLORS["risk"], True, PP_ALIGN.RIGHT)
    add_disclosure(slide, meta)


def page_21(slide, meta, data):
    begin_slide(slide, 21, meta)
    add_circle(slide, 5.35, 2.10, 2.62, COLORS["green"])
    add_text(slide, "完成的\n服務交易", 5.72, 2.86, 1.88, 0.72, 17, COLORS["white"], True, PP_ALIGN.CENTER)
    streams = [
        ("店鋪SaaS", "¥9,800／月", 0.95, 1.55, COLORS["green"]),
        ("預約費", "¥500／單", 9.90, 1.55, COLORS["lime"]),
        ("成功介紹費", "¥30,000／店", 0.95, 4.50, COLORS["yellow"]),
        ("未來CPS結算", "僅合格完單", 9.90, 4.50, COLORS["teal"]),
    ]
    for title, value, x, y, color in streams:
        add_card(slide, x, y, 2.45, 1.02, COLORS["white"], color)
        add_text(slide, title, x + 0.20, y + 0.18, 2.05, 0.28, 11, color, True, PP_ALIGN.CENTER)
        add_text(slide, value, x + 0.20, y + 0.55, 2.05, 0.22, 8.5, COLORS["muted"], True, PP_ALIGN.CENTER)
        line = slide.shapes.add_connector(1, inch(x + 1.23), inch(y + 0.52), inch(6.66), inch(3.41))
        set_shape_line(line, COLORS["line"], 1.1)
    add_pill(slide, "核心模型排除成人性服務；CPS尚無真實歸因收入", 3.30, 5.72, 6.75, 0.46, COLORS["light"], COLORS["risk"], 9.2)
    add_disclosure(slide, meta)


def page_22(slide, meta, data):
    begin_slide(slide, 22, meta)
    econ = data["economics"]
    metrics = [
        ("¥9,800", "SaaS／月", COLORS["green"]),
        ("6個月", "最低合約期", COLORS["lime"]),
        ("¥500", "預約費／單", COLORS["yellow"]),
        ("約¥340", "簡化單筆貢獻｜68%", COLORS["teal"]),
    ]
    for idx, (value, label, color) in enumerate(metrics):
        add_metric(slide, value, label, 0.72 + idx * 3.15, 1.28, 2.80, color)
    add_card(slide, 0.92, 2.85, 5.50, 2.60, COLORS["light"], COLORS["line"])
    add_text(slide, "單筆貢獻算術", 1.28, 3.18, 4.75, 0.30, 13, COLORS["green"], True)
    add_text(slide, "¥500 − 支付3%(¥15) − 風險準備¥30 − 獎勵¥100 − 雲端¥15 ＝ 約¥340", 1.25, 3.72, 4.85, 0.70, 11, COLORS["ink"], True, PP_ALIGN.CENTER)
    add_text(slide, "¥340 ÷ ¥500 ＝ 68%", 1.45, 4.58, 4.45, 0.38, 15, COLORS["teal"], True, PP_ALIGN.CENTER)
    add_card(slide, 6.85, 2.85, 5.55, 2.60, COLORS["white"], COLORS["line"])
    add_text(slide, "兩種渠道的簡化CAC回收", 7.20, 3.18, 4.85, 0.30, 13, COLORS["green"], True)
    paybacks = [("付費獲客", econ["paidStoreCac"], econ["paidStoreCac"]/econ["storeSaasMonthly"], COLORS["green"]), ("介紹成功費", econ["referralSuccessFee"], econ["referralSuccessFee"]/econ["storeSaasMonthly"], COLORS["teal"])]
    for idx, (label, cac, months, color) in enumerate(paybacks):
        y = 3.78 + idx * 0.72
        add_text(slide, f"{label}｜¥{cac:,}", 7.18, y, 2.35, 0.28, 9.5, COLORS["ink"], True)
        bar_width = months * 0.75 + 0.45
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.45, y, bar_width, 0.30, color, color)
        add_text(slide, f"{months:.2f}個月", 9.45, y, bar_width, 0.28, 8.5, COLORS["white"], True, PP_ALIGN.CENTER)
    add_text(slide, "模型算術；預約費免費期仍收SaaS。實際回收需以留存與成本校準。", 7.20, 5.05, 4.85, 0.28, 7.5, COLORS["risk"], True)
    add_disclosure(slide, meta)


def scenario_slide(slide, page, meta, scenario, color, label):
    begin_slide(slide, page, meta)
    categories = ["第1年", "第2年", "第3年"]
    max_value = max(max(scenario["revenueM"]), max(scenario["profitM"]))
    if max_value > 5000:
        minimum, maximum, major = -1000, 10000, 2000
    elif max_value > 1000:
        minimum, maximum, major = -500, 5000, 1000
    else:
        minimum, maximum, major = -200, 800, 200
    add_column_chart(slide, categories, [("營收（百萬日圓）", scenario["revenueM"]), ("營業利益（百萬日圓）", scenario["profitM"])], 0.65, 1.25, 8.10, 4.55, [color, COLORS["dark"]], minimum, maximum, major, False)
    add_card(slide, 9.10, 1.35, 3.40, 4.40, COLORS["light"], COLORS["line"])
    add_text(slide, label, 9.45, 1.68, 2.70, 0.32, 14, color, True, PP_ALIGN.CENTER)
    rows = [
        ("模型店鋪", [fmt_int(v) for v in scenario["stores"]]),
        ("年度訂單", [fmt_int(v) for v in scenario["orders"]]),
        ("營收(M)", [f"{v:.1f}" for v in scenario["revenueM"]]),
        ("營業利益(M)", [f"{v:.1f}" for v in scenario["profitM"]]),
        ("營業利益率", [f"{v:.1f}%" for v in scenario["margin"]]),
    ]
    for ridx, (name, values) in enumerate(rows):
        y = 2.16 + ridx * 0.63
        add_text(slide, name, 9.38, y, 1.02, 0.24, 7.6, COLORS["muted"], True)
        add_text(slide, " → ".join(values), 10.20, y - 0.02, 1.95, 0.35, 7.8, COLORS["ink"], True, PP_ALIGN.RIGHT)
    add_text(slide, "非承諾性模型；需以真實付費、留存、訂單與渠道效率校準。", 9.35, 5.38, 2.95, 0.30, 6.8, COLORS["risk"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_23(slide, meta, data):
    scenario_slide(slide, 23, meta, data["scenarios"]["general"], COLORS["green"], "一般方案")


def page_24(slide, meta, data):
    scenario_slide(slide, 24, meta, data["scenarios"]["aggressive"], "B78837", "激進方案")


def page_25(slide, meta, data):
    begin_slide(slide, 25, meta, page_label="")
    add_doughnut(slide, ["產品與工程", "市場／店鋪導入", "客服／安全／合規", "營運資金"], [80, 60, 30, 30], 0.60, 1.25, 5.45, 4.85, [COLORS["green"], COLORS["lime"], COLORS["yellow"], COLORS["teal"]])
    add_text(slide, "¥200M", 2.05, 2.85, 2.45, 0.55, 25, COLORS["dark"], True, PP_ALIGN.CENTER)
    milestones = [
        ("M0–3", "正式後端、資料、權限與支付／結算設計"),
        ("M4–6", "意向店鋪轉試點；首批付費、完單與留存"),
        ("M7–12", "AI調度實數據；CPS歸因與佣金試點"),
        ("M13–18", "跨店供給池、第二城市與可複製GTM"),
    ]
    add_text(slide, "18個月里程碑", 6.60, 1.42, 5.45, 0.34, 14, COLORS["green"], True)
    for idx, (m, text) in enumerate(milestones):
        y = 2.02 + idx * 0.88
        add_pill(slide, m, 6.65, y, 1.18, 0.44, [COLORS["green"], COLORS["lime"], COLORS["yellow"], COLORS["teal"]][idx], COLORS["white"], 8.8)
        add_text(slide, text, 8.08, y - 0.02, 4.12, 0.48, 9.4, COLORS["ink"], True)
    add_text(slide, "配置為管理層建議，可依投資條款、招募與實際導入節奏調整。", 6.75, 5.55, 5.30, 0.34, 7.8, COLORS["risk"], True)
    add_disclosure(slide, meta)
    add_text(slide, "NeeDo｜海外投資人路演", 0.55, 0.20, 3.0, 0.22, 8.5, COLORS["green"], True)


def page_26(slide, meta, data):
    begin_slide(slide, 26, meta)
    terms = [("¥200M", "本輪融資", COLORS["green"]), ("10%", "出讓股權", COLORS["lime"]), ("¥1.8B", "投前估值", COLORS["yellow"]), ("¥2.0B", "投後估值", COLORS["teal"])]
    for idx, (value, label, color) in enumerate(terms):
        add_metric(slide, value, label, 0.70 + idx * 3.15, 1.25, 2.80, color, "Pre-A" if idx == 0 else None)
    table = slide.shapes.add_table(4, 3, inch(1.00), inch(2.90), inch(7.15), inch(2.25)).table
    headers = ["公司退出股權價值", "投資人10%價值", "對¥200M倍數"]
    rows = [["¥10B", "¥1B", "5×"], ["¥30B", "¥3B", "15×"], ["¥50B", "¥5B", "25×"]]
    for col, header in enumerate(headers):
        table.cell(0, col).text = header
    for row_idx, row in enumerate(rows, 1):
        for col_idx, value in enumerate(row):
            table.cell(row_idx, col_idx).text = value
    for r in range(4):
        for c in range(3):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(COLORS["green"] if r == 0 else COLORS["light"])
            cell.text_frame.margin_left = cell.text_frame.margin_right = inch(0.04)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(8.5 if r == 0 else 10.5)
                    run.font.bold = True
                    run.font.color.rgb = rgb(COLORS["white"] if r == 0 else COLORS["dark"])
    add_card(slide, 8.55, 2.90, 3.70, 2.25, COLORS["light"], COLORS["line"])
    add_text(slide, "只是一組算術情景", 8.90, 3.18, 3.05, 0.34, 14, COLORS["risk"], True, PP_ALIGN.CENTER)
    add_rich_lines(slide, ["假設退出時仍持股10%", "未計後續稀釋與清算優先權", "未計稅務、匯率與退出概率", "不構成估值、退出或回報承諾"], 8.95, 3.72, 2.90, 1.05, 8.4, COLORS["ink"], 3)
    add_disclosure(slide, meta)


def page_27(slide, meta, data):
    scenario_slide(slide, 27, meta, data["scenarios"]["conservative"], COLORS["muted"], "保守情境｜附錄")


def page_28(slide, meta, data):
    begin_slide(slide, 28, meta)
    scenarios = data["scenarios"]
    tracks = [
        ("保守", scenarios["conservative"], COLORS["muted"], "較慢供給拓展／較低訂單密度"),
        ("一般", scenarios["general"], COLORS["green"], "原V2.2高速增長，作為主方案"),
        ("激進", scenarios["aggressive"], "B78837", "更快渠道與資源到位的敏感度"),
    ]
    for idx, (label, scenario, color, note) in enumerate(tracks):
        y = 1.45 + idx * 1.52
        add_pill(slide, label, 0.95, y, 1.10, 0.50, color, COLORS["white"], 10)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 2.35, y, 9.85, 0.50, COLORS["light"], COLORS["light"])
        add_text(slide, f"Y3 {scenario['stores'][2]:,}店｜{scenario['orders'][2]:,}單｜營收 {fmt_m(scenario['revenueM'][2])}｜營業利益率 {scenario['margin'][2]:.1f}%", 2.58, y + 0.07, 6.85, 0.28, 9.4, COLORS["ink"], True)
        add_text(slide, note, 9.45, y + 0.05, 2.45, 0.32, 7.8, color, True, PP_ALIGN.RIGHT)
    add_text(slide, "三種情境的差異來自供給拓展、渠道速度、訂單密度與執行資源；不是三個同時成立的承諾。", 1.35, 5.75, 10.65, 0.42, 13, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_29(slide, meta, data):
    begin_slide(slide, 29, meta)
    scenarios = data["scenarios"]
    categories = [str(i) for i in range(1, 13)]
    add_line_chart(slide, categories, [("保守", scenarios["conservative"]["cashQuarterM"]), ("一般", scenarios["general"]["cashQuarterM"]), ("激進", scenarios["aggressive"]["cashQuarterM"])], 0.72, 1.32, 8.95, 4.60, [COLORS["muted"], COLORS["green"], "B78837"])
    add_text(slide, "季度（1–12）", 4.30, 5.84, 1.75, 0.20, 7.2, COLORS["muted"], False, PP_ALIGN.CENTER)
    add_card(slide, 9.95, 1.45, 2.55, 4.35, COLORS["light"], COLORS["line"])
    add_text(slide, "風險集中窗口", 10.22, 1.83, 2.00, 0.30, 12, COLORS["risk"], True, PP_ALIGN.CENTER)
    add_text(slide, "前10個月", 10.20, 2.35, 2.05, 0.62, 25, COLORS["risk"], True, PP_ALIGN.CENTER)
    add_rich_lines(slide, [f"一般最低現金：¥{scenarios['general']['minCashM']:.1f}M", f"激進最低現金：¥{scenarios['aggressive']['minCashM']:.1f}M", f"保守最低現金：¥{scenarios['conservative']['minCashM']:.1f}M", "轉正月份依情境不同，非保證"], 10.16, 3.18, 2.12, 1.55, 8.5, COLORS["ink"], 3)
    add_disclosure(slide, meta)


def page_30(slide, meta, data):
    begin_slide(slide, 30, meta)
    groups = [
        ("官方市場與監管", data["sources"][:4], COLORS["green"]),
        ("內部財務模型", data["sources"][4:6], COLORS["lime"]),
        ("產品與商談紀錄", data["sources"][6:], COLORS["teal"]),
    ]
    x_positions = [0.72, 4.75, 8.78]
    widths = [3.65, 3.65, 3.82]
    for idx, (title, sources, color) in enumerate(groups):
        x = x_positions[idx]
        add_card(slide, x, 1.32, widths[idx], 4.82, COLORS["white"], color)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, 1.32, widths[idx], 0.65, color, color)
        add_text(slide, title, x + 0.18, 1.50, widths[idx] - 0.36, 0.28, 12, COLORS["white"], True, PP_ALIGN.CENTER)
        y = 2.20
        for source in sources:
            add_text(slide, f"{source['id']}｜{source['title']}", x + 0.25, y, widths[idx] - 0.50, 0.33, 8.7, COLORS["ink"], True)
            add_text(slide, source["url"], x + 0.25, y + 0.34, widths[idx] - 0.50, 0.42, 6.6, COLORS["muted"], False)
            y += 0.92
    add_disclosure(slide, meta)


def page_31(slide, meta, data):
    begin_slide(slide, 31, meta)
    controls = [
        ("資格／業法", "服務類別與資格檢查", COLORS["green"]),
        ("場所規則", "店內／酒店／出張責任邊界", COLORS["lime"]),
        ("隱私與最小權限", "角色化顯示與審計事件", COLORS["yellow"]),
        ("支付／退款", "取消、爭議與結算規則", COLORS["teal"]),
        ("供給責任", "內部、合作與外部供給分層", COLORS["cyan"]),
        ("禁止區隔離", "成人性服務現行禁止", COLORS["blue"]),
    ]
    for idx, (title, body, color) in enumerate(controls):
        row = idx // 3
        col = idx % 3
        x = 0.82 + col * 4.12
        y = 1.35 + row * 2.15
        add_card(slide, x, y, 3.65, 1.62, COLORS["white"], color)
        add_circle(slide, x + 0.20, y + 0.25, 0.48, color)
        add_text(slide, str(idx + 1), x + 0.20, y + 0.35, 0.48, 0.22, 9, COLORS["white"], True, PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.82, y + 0.22, 2.48, 0.30, 11, color, True)
        add_text(slide, body, x + 0.82, y + 0.67, 2.48, 0.55, 8.8, COLORS["ink"], True)
    add_text(slide, "成人市場不是旁支收入，而是與核心模型隔離的受監管未來市場。", 1.45, 5.76, 10.45, 0.38, 13, COLORS["risk"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_32(slide, meta, data):
    begin_slide(slide, 32, meta)
    phases = [
        ("01｜正式產品", "後端、資料庫、RBAC、支付／審計", COLORS["green"]),
        ("02｜商業證據", "意向→試點→付費→留存／完單", COLORS["lime"]),
        ("03｜雙引擎", "AI調度實數據＋CPS歸因試點", COLORS["yellow"]),
        ("04｜城市複製", "跨店供給池、第二城市、可複製GTM", COLORS["teal"]),
    ]
    for idx, (title, body, color) in enumerate(phases):
        y = 1.42 + idx * 1.08
        add_pill(slide, title, 0.95, y, 2.35, 0.56, color, COLORS["white"], 10)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, 3.45, y, 8.55, 0.56, COLORS["light"], COLORS["light"])
        add_text(slide, body, 3.78, y + 0.10, 7.75, 0.30, 10, COLORS["ink"], True)
        if idx < 3:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.DOWN_ARROW, 6.40, y + 0.63, 0.50, 0.32, COLORS["beige"], COLORS["beige"])
    add_pill(slide, "18個月：把可操作原型與使用意向轉成可驗證商業資料", 2.00, 5.88, 9.35, 0.46, COLORS["dark"], COLORS["white"], 11)
    add_disclosure(slide, meta)


def page_33(slide, meta, data):
    begin_slide(slide, 33, meta)
    folders = [
        ("商業", ["CRM意向清單", "試點／合約／收款", "城市進場假設"], COLORS["green"]),
        ("產品", ["可操作原型", "架構／權限／審計", "AI／CPS事件設計"], COLORS["lime"]),
        ("財務", ["V2.2模型", "一般／激進／保守", "單位經濟校準"], COLORS["yellow"]),
        ("合規", ["資格與業法", "隱私／支付／退款", "禁止業務隔離"], COLORS["teal"]),
    ]
    for idx, (title, items, color) in enumerate(folders):
        x = 0.80 + idx * 3.12
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.FOLDED_CORNER, x, 1.40, 2.70, 3.95, color, color)
        add_text(slide, title, x + 0.30, 1.76, 2.10, 0.34, 15, COLORS["white"], True, PP_ALIGN.CENTER)
        for item_idx, item in enumerate(items):
            add_pill(slide, item, x + 0.30, 2.55 + item_idx * 0.70, 2.10, 0.45, COLORS["white"], color, 8.2)
    add_text(slide, "我們主動把未驗證項目攤開：用資料室、合約、收款、留存與事件資料逐項關閉。", 1.30, 5.72, 10.75, 0.45, 13, COLORS["dark"], True, PP_ALIGN.CENTER)
    add_disclosure(slide, meta)


def page_34(slide, meta, data):
    remove_negative_position_shapes(slide)
    clear_text_recursive(slide.shapes)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 3.15, 1.15, 5.25, 2.45, COLORS["white"], COLORS["white"])
    add_text(slide, "NeeDo", 3.42, 1.30, 4.70, 0.52, 30, COLORS["green"], True, PP_ALIGN.CENTER)
    add_text(slide, "讓服務被看見、被選擇、被完成、被分配", 3.40, 1.92, 4.74, 0.86, 17, COLORS["ink"], True, PP_ALIGN.CENTER)
    add_text(slide, "Pre-A｜融資 2億日圓｜出讓 10% 股權", 3.42, 2.90, 4.70, 0.30, 12.5, COLORS["green"], True, PP_ALIGN.CENTER)
    next_steps = ["產品／資料室展示", "試點店鋪證據核對", "模型與敏感度討論", "條款與18個月里程碑"]
    for idx, item in enumerate(next_steps):
        add_pill(slide, f"{idx+1:02d}  {item}", 1.05 + idx * 3.05, 5.75, 2.72, 0.50, [COLORS["green"], COLORS["lime"], COLORS["yellow"], COLORS["teal"]][idx], COLORS["white"], 8.5)
    add_text(slide, "投資人討論用｜2026年8月", 4.85, 6.55, 3.70, 0.22, 8, COLORS["white"], False, PP_ALIGN.CENTER)


PAGE_BUILDERS = {
    1: page_1, 2: page_2, 3: page_3, 4: page_4, 5: page_5, 6: page_6, 7: page_7, 8: page_8, 9: page_9,
    10: page_10, 11: page_11, 12: page_12, 13: page_13, 14: page_14, 15: page_15, 16: page_16, 17: page_17, 18: page_18,
    19: page_19, 20: page_20, 21: page_21, 22: page_22, 23: page_23, 24: page_24, 25: page_25, 26: page_26,
    27: page_27, 28: page_28, 29: page_29, 30: page_30, 31: page_31, 32: page_32, 33: page_33, 34: page_34,
}


def build(template: Path, output: Path) -> None:
    data = load_payload()
    slides_meta = data["premiumSlides"]
    if len(slides_meta) != 34:
        raise ValueError(f"Expected 34 approved slides, got {len(slides_meta)}")
    prs = Presentation(template)
    set_slide_order(prs, TEMPLATE_PAGE_MAP)
    if len(prs.slides) != 34:
        raise ValueError(f"Expected 34 template-mapped slides, got {len(prs.slides)}")
    for page, slide in enumerate(prs.slides, start=1):
        meta = slides_meta[page - 1]
        if meta["page"] != page:
            raise ValueError(f"Content page mismatch at {page}")
        if page not in (1, 34):
            remove_all_shapes(slide)
        PAGE_BUILDERS[page](slide, meta, data)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(json.dumps({"output": str(output), "slides": len(prs.slides), "template": str(template)}, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.template.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
