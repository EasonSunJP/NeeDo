#!/usr/bin/env python3
from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_PPTX = ROOT / "outputs/needo-roadshow-template-2026-08-23/NeeDo_海外投資人路演_BP_繁中_模板版_2026-08-23.pptx"
OUTPUT_DIR = DEFAULT_PPTX.parent


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def main(path: Path) -> None:
    prs = Presentation(path)
    errors = []
    if len(prs.slides) != 34:
        errors.append(f"expected 34 slides, got {len(prs.slides)}")
    ratio = prs.slide_width / prs.slide_height
    if abs(ratio - 16 / 9) > 0.001:
        errors.append(f"expected 16:9, got {ratio}")

    all_text = []
    chart_count = 0
    bounds_errors = []
    for page, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                chart_count += 1
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    all_text.append(f"[{page:02d}] {text}")
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 12700 or shape.top + shape.height > prs.slide_height + 12700:
                bounds_errors.append({"page": page, "name": shape.name, "left": shape.left, "top": shape.top, "width": shape.width, "height": shape.height})
    if chart_count != 5:
        errors.append(f"expected 5 native charts, got {chart_count}")
    if bounds_errors:
        errors.append(f"out-of-bounds top-level shapes: {len(bounds_errors)}")

    joined = "\n".join(all_text)
    required = [
        "Pre-A｜融資 2億日圓｜出讓 10% 股權",
        "100+",
        "不等於已簽約",
        "不等於已付費／活躍店鋪",
        "不等於收入或GMV",
        "¥1.8B",
        "¥2.0B",
        "一般方案",
        "激進方案",
        "保守情境｜附錄",
        "成人性服務現行禁止",
        "尚無真實歸因 GMV 或 CPS 收入",
        "不構成估值、退出或回報承諾",
    ]
    missing = [item for item in required if item not in joined]
    if missing:
        errors.append(f"missing required text: {missing}")

    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            errors.append(f"bad zip entry: {bad}")
        theme = archive.read("ppt/theme/theme1.xml").decode("utf-8")
        for color in ["549E39", "8AB833", "C0CF3A", "029676", "4AB5C4", "0989B1"]:
            if color not in theme:
                errors.append(f"template theme color missing: {color}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "extracted-text-template.txt").write_text(joined, encoding="utf-8")
    report = {
        "pptx": str(path),
        "slides": len(prs.slides),
        "ratio": ratio,
        "nativeCharts": chart_count,
        "outOfBounds": bounds_errors,
        "requiredTextMissing": missing,
        "errors": errors,
        "status": "passed" if not errors else "failed",
    }
    (OUTPUT_DIR / "verification-template.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main(Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_PPTX)
